from pathlib import Path
import shutil
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

HEADERS = [
    "ID документа", "Путь к документу", "Имя файла", "Формат", "Статус", "Дата обработки",
    "Номер документа", "Дата документа", "Организация", "ИНН", "Номер договора", "Тип документа",
    "Отправитель", "Получатель", "Тема письма", "Краткое содержание", "Требуемые действия",
    "Срок исполнения", "Нужен ответ", "Уверенность", "Требуется проверка", "Ошибка"
]
FIELD_MAP = {"Номер документа":"document_number", "Дата документа":"document_date",
             "Организация":"organization", "ИНН":"inn", "Номер договора":"contract_number",
             "Тип документа":"document_type"}

class ExcelRegistryWriter:
    def append_results(self, output_path, records, create_if_missing=True, backup=True):
        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        if path.exists():
            if backup:
                stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
                shutil.copy2(path, path.with_name(f"{path.stem}.{stamp}.bak{path.suffix}"))
            wb = load_workbook(path)
            ws = wb["Реестр"] if "Реестр" in wb.sheetnames else wb.active
            if ws.title != "Реестр" and ws.max_row <= 1:
                ws.title = "Реестр"
            headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
            if not headers or all(v is None for v in headers):
                ws.delete_rows(1, ws.max_row)
                headers = []
            for header in HEADERS:
                if header not in headers:
                    headers.append(header)
                    ws.cell(1, len(headers), header)
        else:
            if not create_if_missing:
                raise FileNotFoundError(path)
            wb = Workbook()
            ws = wb.active
            ws.title = "Реестр"
            headers = HEADERS[:]
            ws.append(headers)
        self._style_header(ws)
        path_col = headers.index("Путь к документу") + 1
        existing_rows = {}
        for row_num in range(2, ws.max_row + 1):
            key = ws.cell(row_num, path_col).value
            if key:
                existing_rows[str(Path(str(key)).resolve()).casefold()] = row_num
        added = 0
        for record in records:
            file_path = str(Path(record.get("file_path", "")).resolve())
            key = file_path.casefold()
            row_num = existing_rows.get(key)
            if row_num is None:
                row_num = ws.max_row + 1
                existing_rows[key] = row_num
                added += 1
            fields = record.get("fields", {})
            data = {"ID документа": record.get("sha256", ""), "Путь к документу": file_path,
                    "Имя файла": record.get("file_name", ""), "Формат": record.get("file_type", ""),
                    "Статус": record.get("status", ""), "Дата обработки": record.get("processed_at", ""),
                    "Краткое содержание": record.get("summary", ""), "Требуемые действия": "\n".join(record.get("required_actions", [])),
                    "Срок исполнения": record.get("deadline", ""), "Нужен ответ": "Да" if record.get("needs_reply") else "Не определено",
                    "Уверенность": record.get("confidence", ""), "Требуется проверка": "Да" if record.get("review_required") else "Нет",
                    "Ошибка": record.get("error", ""), "Отправитель": record.get("sender", ""),
                    "Получатель": record.get("recipient", ""), "Тема письма": record.get("subject", "")}
            for header, field_name in FIELD_MAP.items():
                data[header] = fields.get(field_name, record.get(field_name, ""))
            for col, header in enumerate(headers, 1):
                if header in data:
                    ws.cell(row_num, col, data[header])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for column in ws.columns:
            letter = column[0].column_letter
            max_len = min(max((len(str(c.value or "")) for c in column[:100]), default=10), 48)
            ws.column_dimensions[letter].width = max(12, max_len + 2)
        # Full text is stored separately to keep the main register readable.
        full_ws = wb["Полные тексты"] if "Полные тексты" in wb.sheetnames else wb.create_sheet("Полные тексты")
        if full_ws.max_row == 1 and full_ws.cell(1, 1).value is None:
            for c, header in enumerate(["ID документа", "Путь к документу", "Имя файла", "Полный текст", "Метаданные письма", "Вложения", "Предупреждения извлечения"], 1):
                full_ws.cell(1, c, header)
        full_headers = [full_ws.cell(1, c).value for c in range(1, full_ws.max_column + 1)]
        full_path_col = full_headers.index("Путь к документу") + 1
        full_rows = {str(full_ws.cell(r, full_path_col).value).casefold(): r for r in range(2, full_ws.max_row + 1) if full_ws.cell(r, full_path_col).value}
        import json
        for record in records:
            p = str(Path(record.get("file_path", "")).resolve())
            r = full_rows.get(p.casefold())
            if r is None:
                r = full_ws.max_row + 1
                full_rows[p.casefold()] = r
            values = [record.get("sha256", ""), p, record.get("file_name", ""), record.get("full_text", ""),
                      json.dumps(record.get("metadata", {}), ensure_ascii=False),
                      json.dumps(record.get("attachments", []), ensure_ascii=False),
                      "\n".join(record.get("extraction_warnings", []))]
            for c, value in enumerate(values, 1): full_ws.cell(r, c, value)
            full_ws.cell(r, 4).alignment = Alignment(wrap_text=True, vertical="top")
        full_ws.freeze_panes = "A2"
        full_ws.column_dimensions["D"].width = 100
        full_ws.column_dimensions["D"].alignment = Alignment(wrap_text=True, vertical="top")
        self._style_header(full_ws)
        attachments_ws = wb["Вложения"] if "Вложения" in wb.sheetnames else wb.create_sheet("Вложения")
        if attachments_ws.max_row == 1 and attachments_ws.cell(1, 1).value is None:
            for c, header in enumerate(["ID письма", "Путь к письму", "Имя вложения", "Тип содержимого", "Статус обработки", "Путь к сохранённому вложению"], 1):
                attachments_ws.cell(1, c, header)
        # Rebuild managed attachment metadata from the current export batch to avoid duplicate rows on reruns.
        if attachments_ws.max_row > 1:
            attachments_ws.delete_rows(2, attachments_ws.max_row - 1)
        for record in records:
            for attachment in record.get("attachments", []):
                attachments_ws.append([record.get("sha256", ""), record.get("file_path", ""), attachment.get("filename", "attachment"), attachment.get("content_type", ""), "Обнаружено; содержимое не извлечено", ""])
        self._style_header(attachments_ws)
        wb.save(path)
        wb.close()
        return added

    @staticmethod
    def _style_header(ws):
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="243B53")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[1].height = 32
