from dataclasses import dataclass, field
from pathlib import Path
from datetime import datetime, timezone
from email import policy
from email.parser import BytesParser
from email.header import decode_header, make_header
import logging
import re

logger = logging.getLogger(__name__)

@dataclass
class DocumentData:
    file_path: str
    file_name: str
    file_type: str
    file_size: int
    modified_at: str
    status: str
    text: str = ""
    error: str = ""
    metadata: dict = field(default_factory=dict)
    extraction_warnings: list = field(default_factory=list)
    attachments: list = field(default_factory=list)

    @property
    def text_length(self):
        return len(self.text)

class DocumentService:
    def __init__(self, ocr_enabled=True, ocr_language="rus+eng", ocr_min_chars=40):
        self.ocr_enabled = ocr_enabled
        self.ocr_language = ocr_language
        self.ocr_min_chars = ocr_min_chars

    def read(self, path):
        path = Path(path)
        stat = path.stat()
        base = dict(file_path=str(path.resolve()), file_name=path.name,
                    file_type=path.suffix.lower().lstrip("."), file_size=stat.st_size,
                    modified_at=datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat())
        try:
            suffix = path.suffix.lower()
            metadata, warnings, attachments = {}, [], []
            if suffix == ".pdf":
                text, warnings = self._read_pdf(path)
            elif suffix == ".docx":
                text = self._read_docx(path)
            elif suffix == ".xlsx":
                text = self._read_xlsx(path)
            elif suffix == ".eml":
                text, metadata, attachments = self._read_eml(path)
            elif suffix == ".msg":
                text, metadata, attachments = self._read_msg(path)
            else:
                raise ValueError(f"Формат не поддерживается: {suffix}")
            return DocumentData(**base, status="ok", text=text, metadata=metadata,
                                extraction_warnings=warnings, attachments=attachments)
        except Exception as exc:
            logger.exception("Не удалось прочитать %s", path)
            return DocumentData(**base, status="error", error=f"{type(exc).__name__}: {exc}")

    def _read_pdf(self, path):
        try:
            import fitz
        except ImportError as exc:
            raise RuntimeError("Не установлен PyMuPDF. Установите зависимости из requirements.txt.") from exc
        chunks, warnings = [], []
        with fitz.open(path) as doc:
            for index, page in enumerate(doc, 1):
                text = page.get_text("text") or ""
                if len(text.strip()) < self.ocr_min_chars and self.ocr_enabled:
                    try:
                        pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                        import pytesseract
                        from PIL import Image
                        import io
                        image = Image.open(io.BytesIO(pix.tobytes("png")))
                        ocr_text = pytesseract.image_to_string(image, lang=self.ocr_language)
                        if len(ocr_text.strip()) > len(text.strip()):
                            text = ocr_text
                    except ImportError:
                        warnings.append(f"Страница {index}: OCR недоступен (установите pytesseract и Pillow).")
                    except Exception as exc:
                        warnings.append(f"Страница {index}: OCR не выполнен: {exc}")
                chunks.append(f"[Страница {index}]\n{text}".strip())
        return "\n\n".join(chunks), warnings

    @staticmethod
    def _read_docx(path):
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("Не установлен python-docx. Установите зависимости из requirements.txt.") from exc
        doc = Document(path)
        chunks = [p.text for p in doc.paragraphs if p.text]
        for table in doc.tables:
            for row in table.rows:
                chunks.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(chunks)

    @staticmethod
    def _read_xlsx(path):
        from openpyxl import load_workbook
        chunks = []
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                chunks.append(f"[Лист: {ws.title}]")
                for row in ws.iter_rows(values_only=True):
                    values = [str(v) for v in row if v is not None]
                    if values:
                        chunks.append(" | ".join(values))
        finally:
            wb.close()
        return "\n".join(chunks)

    @staticmethod
    def _decode_header(value):
        if not value:
            return ""
        try:
            return str(make_header(decode_header(str(value))))
        except Exception:
            return str(value)

    @classmethod
    def _read_eml(cls, path):
        with path.open("rb") as f:
            msg = BytesParser(policy=policy.default).parse(f)
        metadata = {"sender": cls._decode_header(msg.get("From")),
                    "recipient": cls._decode_header(msg.get("To")),
                    "cc": cls._decode_header(msg.get("Cc")),
                    "subject": cls._decode_header(msg.get("Subject")),
                    "email_date": cls._decode_header(msg.get("Date")),
                    "message_id": str(msg.get("Message-ID", ""))}
        bodies, attachments = [], []
        if msg.is_multipart():
            for part in msg.walk():
                if part.is_multipart():
                    continue
                filename = part.get_filename()
                if filename or part.get_content_disposition() == "attachment":
                    attachments.append({"filename": cls._decode_header(filename) or "attachment", "content_type": part.get_content_type()})
                    continue
                if part.get_content_type() in ("text/plain", "text/html"):
                    try:
                        body = part.get_content()
                    except Exception:
                        payload = part.get_payload(decode=True) or b""
                        body = payload.decode(part.get_content_charset() or "utf-8", errors="replace")
                    if part.get_content_type() == "text/html":
                        body = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", body)
                        body = re.sub(r"(?s)<[^>]+>", " ", body)
                        body = re.sub(r"&nbsp;", " ", body)
                    if body and body.strip(): bodies.append(body.strip())
        else:
            try: body = msg.get_content()
            except Exception: body = (msg.get_payload(decode=True) or b"").decode("utf-8", errors="replace")
            if msg.get_content_type() == "text/html":
                body = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", body)
                body = re.sub(r"(?s)<[^>]+>", " ", body)
            bodies.append(str(body).strip())
        header_text = "\n".join(f"{k}: {v}" for k, v in metadata.items() if v)
        return (header_text + "\n\n" + "\n\n".join(bodies)).strip(), metadata, attachments

    @staticmethod
    def _read_msg(path):
        try:
            import extract_msg
        except ImportError as exc:
            raise RuntimeError("Для MSG установите extract-msg из requirements.txt.") from exc
        msg = extract_msg.Message(str(path))
        try:
            metadata = {"sender": str(getattr(msg, "sender", "") or ""),
                        "recipient": str(getattr(msg, "to", "") or ""),
                        "cc": str(getattr(msg, "cc", "") or ""),
                        "subject": str(getattr(msg, "subject", "") or ""),
                        "email_date": str(getattr(msg, "date", "") or ""),
                        "message_id": str(getattr(msg, "messageId", "") or "")}
            body = str(getattr(msg, "body", "") or getattr(msg, "htmlBody", "") or "")
            if "<" in body and ">" in body:
                body = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", body)
                body = re.sub(r"(?s)<[^>]+>", " ", body)
            header_text = "\n".join(f"{k}: {v}" for k, v in metadata.items() if v)
            attachments = [{"filename": str(getattr(a, "longFilename", None) or getattr(a, "shortFilename", None) or "attachment"),
                            "content_type": "unknown"} for a in (getattr(msg, "attachments", None) or [])]
            return (header_text + "\n\n" + body).strip(), metadata, attachments
        finally:
            try: msg.close()
            except Exception: pass
