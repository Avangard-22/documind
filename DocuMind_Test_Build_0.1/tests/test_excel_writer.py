import tempfile, unittest
from pathlib import Path
from openpyxl import load_workbook
from app.excel_writer import ExcelRegistryWriter

class ExcelWriterTests(unittest.TestCase):
    def test_create_and_update_existing_path(self):
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/"out.xlsx"; writer=ExcelRegistryWriter()
            record={"file_path":str(Path(d)/"doc.pdf"),"file_name":"doc.pdf","file_type":"pdf","status":"ok","processed_at":"now","fields":{"document_number":"N-1"},"summary":"test"}
            self.assertEqual(writer.append_results(path,[record]),1)
            record["fields"]["document_number"]="N-2"
            self.assertEqual(writer.append_results(path,[record]),0)
            wb=load_workbook(path,data_only=True); ws=wb.active
            headers=[c.value for c in ws[1]]
            row=[c.value for c in ws[2]]
            self.assertEqual(ws.max_row,2)
            self.assertEqual(row[headers.index("Номер документа")],"N-2")
            wb.close()

if __name__=="__main__":unittest.main()
